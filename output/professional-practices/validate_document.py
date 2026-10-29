from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from collections import Counter
p=Path(r'E:\Targeted Project\bookforge\output\professional-practices\Professional_Practices_Combined.docx')
with ZipFile(p) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith('.xml'): etree.fromstring(z.read(name))
d=Document(p)
text='\n'.join(x.text for x in d.paragraphs)
for term in ['Presentation Drawing','Working Drawing','Detail Drawing','Shop Drawing','As-Built Drawing','মূলনীতি ১','মূলনীতি ২','মূলনীতি ৩','মূলনীতি ৪','মূলনীতি ৫','মূলনীতি ৬','মূলনীতি ৭','মূলনীতি ৮','১৪ দিনের','২.৭.৫','প্রবেশগম্যতা','আপিল','সংকলনের উৎস']:
 assert term in text, term
assert '\ufffd' not in text
assert not d.element.xpath('//w:del | //w:ins')
counts=Counter(x.text.strip() for x in d.paragraphs if x.text.strip())
assert not [t for t,n in counts.items() if n>1]
print('PASS: DOCX ZIP integrity and XML parsing')
print('PASS: all five drawing types, eight principles, complaints process and source references present')
print('PASS: valid Bangla text; no replacement characters, duplicate paragraphs or tracked changes')
print('Final file size:',p.stat().st_size,'bytes')
